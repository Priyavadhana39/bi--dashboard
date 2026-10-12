"""
Document Loader for RAG System
Loads procurement policy documents from various formats.
"""

import os
from typing import List, Dict, Any
from pathlib import Path
import logging

logger = logging.getLogger(__name__)


class Document:
    """Represents a loaded document"""
    
    def __init__(self, content: str, metadata: Dict[str, Any]):
        self.content = content
        self.metadata = metadata
    
    def __repr__(self):
        return f"Document(source={self.metadata.get('source', 'unknown')}, length={len(self.content)})"


class DocumentLoader:
    """Load documents from various formats"""
    
    def __init__(self, knowledge_base_path: str):
        self.knowledge_base_path = Path(knowledge_base_path)
        
    def load_documents(self) -> List[Document]:
        """Load all documents from the knowledge base directory"""
        documents = []
        
        if not self.knowledge_base_path.exists():
            logger.warning(f"Knowledge base path does not exist: {self.knowledge_base_path}")
            return documents
        
        # Load text files
        for txt_file in self.knowledge_base_path.glob('**/*.txt'):
            try:
                doc = self.load_text_file(txt_file)
                if doc:
                    documents.append(doc)
            except Exception as e:
                logger.error(f"Failed to load {txt_file}: {e}")
        
        # Load markdown files
        for md_file in self.knowledge_base_path.glob('**/*.md'):
            try:
                doc = self.load_markdown_file(md_file)
                if doc:
                    documents.append(doc)
            except Exception as e:
                logger.error(f"Failed to load {md_file}: {e}")
        
        # Load PDF files
        for pdf_file in self.knowledge_base_path.glob('**/*.pdf'):
            try:
                doc = self.load_pdf_file(pdf_file)
                if doc:
                    documents.append(doc)
            except Exception as e:
                logger.error(f"Failed to load {pdf_file}: {e}")
        
        # Load DOCX files
        for docx_file in self.knowledge_base_path.glob('**/*.docx'):
            try:
                doc = self.load_docx_file(docx_file)
                if doc:
                    documents.append(doc)
            except Exception as e:
                logger.error(f"Failed to load {docx_file}: {e}")
        
        logger.info(f"Loaded {len(documents)} documents from {self.knowledge_base_path}")
        return documents
    
    def load_text_file(self, file_path: Path) -> Document:
        """Load a plain text file"""
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        metadata = {
            'source': str(file_path.name),
            'file_path': str(file_path),
            'file_type': 'text',
            'file_size': file_path.stat().st_size
        }
        
        return Document(content=content, metadata=metadata)
    
    def load_markdown_file(self, file_path: Path) -> Document:
        """Load a markdown file"""
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        metadata = {
            'source': str(file_path.name),
            'file_path': str(file_path),
            'file_type': 'markdown',
            'file_size': file_path.stat().st_size
        }
        
        return Document(content=content, metadata=metadata)
    
    def load_pdf_file(self, file_path: Path) -> Document:
        """Load a PDF file"""
        try:
            from pypdf import PdfReader
            
            reader = PdfReader(file_path)
            content = ""
            
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    content += f"\n--- Page {page_num + 1} ---\n"
                    content += text
            
            metadata = {
                'source': str(file_path.name),
                'file_path': str(file_path),
                'file_type': 'pdf',
                'file_size': file_path.stat().st_size,
                'num_pages': len(reader.pages)
            }
            
            return Document(content=content, metadata=metadata)
            
        except ImportError:
            logger.warning("pypdf not installed, skipping PDF file")
            return None
        except Exception as e:
            logger.error(f"Error loading PDF {file_path}: {e}")
            return None
    
    def load_docx_file(self, file_path: Path) -> Document:
        """Load a DOCX file"""
        try:
            from docx import Document as DocxDocument
            
            doc = DocxDocument(file_path)
            content = "\n\n".join([paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()])
            
            metadata = {
                'source': str(file_path.name),
                'file_path': str(file_path),
                'file_type': 'docx',
                'file_size': file_path.stat().st_size,
                'num_paragraphs': len(doc.paragraphs)
            }
            
            return Document(content=content, metadata=metadata)
            
        except ImportError:
            logger.warning("python-docx not installed, skipping DOCX file")
            return None
        except Exception as e:
            logger.error(f"Error loading DOCX {file_path}: {e}")
            return None
